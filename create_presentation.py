import sys
sys.path.insert(0, '/tmp/pptlib')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR
from pathlib import Path

OUT=Path('presentations/backlog-buddy-overview.pptx'); OUT.parent.mkdir(exist_ok=True)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BG=RGBColor(248,249,253); WHITE=RGBColor(255,255,255); INK=RGBColor(24,35,59); MUTED=RGBColor(102,113,134)
PURPLE=RGBColor(105,87,232); LAV=RGBColor(235,232,255); NAVY=RGBColor(32,42,79); MINT=RGBColor(223,246,237); GREEN=RGBColor(31,148,105)
CORAL=RGBColor(255,232,227); ORANGE=RGBColor(210,91,66); GOLD=RGBColor(255,241,204); BLUE=RGBColor(228,239,255)

def rect(slide,x,y,w,h,fill=WHITE,radius=True,line=None):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=fill; shp.line.color.rgb=line or fill
    if radius: shp.adjustments[0]=0.12
    return shp

def text(slide,txt,x,y,w,h,size=18,color=INK,bold=False,font='Aptos',align=PP_ALIGN.LEFT,margin=.04):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.text=txt; p.alignment=align
    r=p.runs[0]; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def rich(slide,items,x,y,w,h,size=18):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.04); tf.margin_top=tf.margin_bottom=Inches(.03)
    for i,(title,body,color) in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(12)
        r=p.add_run(); r.text=title; r.font.name='Aptos'; r.font.bold=True; r.font.size=Pt(size); r.font.color.rgb=color
        r=p.add_run(); r.text=body; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=INK
    return box

def base(title=None,kicker=None,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); bg=s.background.fill; bg.solid(); bg.fore_color.rgb=NAVY if dark else BG
    if kicker: text(s,kicker.upper(),.7,.35,6,.3,10,PURPLE if not dark else RGBColor(168,154,255),True)
    if title: text(s,title,.7,.72,11.9,.7,29,WHITE if dark else INK,True)
    return s

def footer(s,n,dark=False):
    c=RGBColor(184,190,205) if dark else MUTED
    text(s,'BACKLOG BUDDY  •  PREPARE SMART. CLEAR YOUR BACKLOGS.',.7,7.12,8,.2,8,c,True)
    text(s,f'{n:02d}',12.1,7.08,.5,.25,9,c,True,align=PP_ALIGN.RIGHT)

def pill(s,label,x,y,w,fill=LAV,color=PURPLE):
    rect(s,x,y,w,.36,fill); text(s,label,x,y,w,.36,10,color,True,align=PP_ALIGN.CENTER)

def icon_badge(s,label,x,y,fill=LAV,color=PURPLE,size=.55):
    rect(s,x,y,size,size,fill); text(s,label,x,y,size,size,13,color,True,align=PP_ALIGN.CENTER)

# 1 Title
s=base(dark=True)
rect(s,.75,.65,.62,.62,PURPLE); text(s,'BB',.75,.65,.62,.62,15,WHITE,True,align=PP_ALIGN.CENTER)
text(s,'BACKLOG BUDDY',1.52,.7,3.1,.4,13,RGBColor(184,176,255),True)
text(s,'Prepare smart.\nMake your comeback.',.75,1.65,7.4,1.7,44,WHITE,True)
text(s,'A focused engineering backlog preparation platform',.78,3.58,6.5,.45,20,RGBColor(198,204,220))
for i,(lab,w) in enumerate([('SUBJECT-WISE',1.45),('MOBILE-FIRST',1.35),('STUDENT-FRIENDLY',1.7)]): pill(s,lab,.78+sum([1.6,1.5,1.85][:i]),4.42,w,RGBColor(49,59,103),RGBColor(215,211,255))
# visual card
rect(s,8.55,1.05,3.95,5.15,WHITE); pill(s,'NEXT EXAM',8.92,1.42,1.05,GOLD,RGBColor(144,97,24)); text(s,'Database Management\nSystems',8.92,1.95,3.05,.8,22,INK,True)
text(s,'Preparation progress',8.92,3.0,2.2,.3,11,MUTED); text(s,'68%',11.55,3.0,.55,.3,11,PURPLE,True,align=PP_ALIGN.RIGHT)
rect(s,8.92,3.43,3.15,.11,RGBColor(230,233,239)); rect(s,8.92,3.43,2.14,.11,PURPLE)
for i,(v,l) in enumerate([('42','Questions'),('8','Papers'),('5','Units')]):
    text(s,v,8.9+i*1.08,4.08,1,.35,21,INK,True,align=PP_ALIGN.CENTER); text(s,l,8.9+i*1.08,4.43,1,.25,9,MUTED,align=PP_ALIGN.CENTER)
rect(s,8.92,5.18,3.15,.56,NAVY); text(s,'CONTINUE PREPARING  →',8.92,5.18,3.15,.56,11,WHITE,True,align=PP_ALIGN.CENTER)
footer(s,1,True)

# 2 Problem
s=base('The problem students face','THE CHALLENGE')
text(s,'Backlog preparation is often fragmented, stressful and difficult to organize.',.7,1.5,7.6,.58,22,MUTED)
cards=[('01','Resources everywhere','Students search across groups, drives and websites.',LAV,PURPLE),('02','No clear priority','Important topics and repeated patterns are hard to spot.',CORAL,ORANGE),('03','Limited visibility','Progress, deadlines and revision plans stay disconnected.',MINT,GREEN)]
for i,(n,t,b,f,c) in enumerate(cards):
    x=.7+i*4.08; rect(s,x,2.45,3.76,3.25,WHITE); icon_badge(s,n,x+.28,2.76,f,c,.62); text(s,t,x+.28,3.62,3.1,.42,19,INK,True); text(s,b,x+.28,4.18,3.05,.85,14,MUTED)
footer(s,2)

# 3 Solution
s=base('One calm place to prepare','THE SOLUTION')
text(s,'Backlog Buddy organizes preparation around the student’s actual curriculum.',.7,1.45,8.4,.45,20,MUTED)
steps=['University','Course','Branch','Regulation','Year','Semester','Subject']
for i,st in enumerate(steps):
    x=.7+i*1.76; rect(s,x,2.35,1.48,.82,LAV if i<6 else PURPLE); text(s,f'{i+1}',x+.1,2.49,.28,.35,12,PURPLE if i<6 else WHITE,True,align=PP_ALIGN.CENTER); text(s,st,x+.38,2.39,1.0,.5,10,INK if i<6 else WHITE,True,align=PP_ALIGN.CENTER)
    if i<6:text(s,'→',x+1.5,2.5,.25,.25,14,MUTED,True,align=PP_ALIGN.CENTER)
rect(s,.7,3.75,11.92,2.2,WHITE)
items=[('Personalized results  ','Only the right subjects, units and resources.  ',PURPLE),('Faster discovery  ','Search questions, topics, papers and notes.  ',GREEN),('Actionable preparation  ','Practice, plan, save and track progress.',ORANGE)]
rich(s,items,1.05,4.08,11.2,1.45,17)
footer(s,3)

# 4 Features
s=base('Everything needed for focused preparation','CORE FEATURES')
features=[('PQ','Previous Papers','Searchable, filterable and license-aware.',BLUE,RGBColor(49,112,187)),('IQ','Important Questions','Organized by unit, marks and frequency.',CORAL,ORANGE),('MP','Model Papers','Original downloadable practice PDFs.',MINT,GREEN),('SY','Syllabus','Unit structure and preparation scope.',GOLD,RGBColor(162,111,22)),('PR','Practice','MCQs, explanations, timers and history.',LAV,PURPLE),('PL','Study Planner','Balanced timetable based on exam date.',BLUE,RGBColor(49,112,187))]
for i,(ic,t,b,f,c) in enumerate(features):
    col=i%3; row=i//3; x=.7+col*4.08; y=1.55+row*2.45
    rect(s,x,y,3.76,2.08,WHITE); icon_badge(s,ic,x+.25,y+.25,f,c,.55); text(s,t,x+.98,y+.25,2.4,.38,16,INK,True); text(s,b,x+.25,y+1.02,3.15,.62,12,MUTED)
footer(s,4)

# 5 Student tools
s=base('Built around daily student habits','STUDENT EXPERIENCE')
# left dashboard
rect(s,.7,1.45,7.4,4.95,WHITE); text(s,'MY DASHBOARD',1.03,1.78,2,.25,10,PURPLE,True)
for i,(num,lab,clr) in enumerate([('4','My subjects',LAV),('3','My backlogs',CORAL),('68%','Practice progress',MINT)]):
    x=1.03+i*2.23; rect(s,x,2.28,2.0,1.18,clr); text(s,num,x+.16,2.42,1.65,.4,22,INK,True); text(s,lab,x+.16,2.87,1.65,.25,10,MUTED)
text(s,'TODAY’S PLAN',1.03,3.86,2,.25,10,MUTED,True)
for i,(a,b,c) in enumerate([('✓','DBMS · Normalization','Completed'),('○','OS · Deadlocks','Up next'),('○','CN · Transport Layer','45 min')]):
    y=4.23+i*.57; text(s,a,1.03,y,.3,.3,15,GREEN if i==0 else MUTED,True); text(s,b,1.42,y,3.5,.28,12,INK,True); text(s,c,6.33,y,.95,.28,10,MUTED,align=PP_ALIGN.RIGHT)
# right benefits
text(s,'STUDENT TOOLS',8.65,1.63,2.5,.3,10,PURPLE,True)
items=[('Save  ','Questions and complete filtered lists',PURPLE),('Track  ','Practice scores, history and progress',GREEN),('Plan  ','Daily study blocks and revision time',ORANGE),('Install  ','Use as a mobile app with offline access',PURPLE)]
rich(s,items,8.65,2.15,3.55,3.5,15)
footer(s,5)

# 6 Practice
s=base('Practice that explains, not just scores','ACTIVE RECALL')
rect(s,.7,1.5,8.2,4.95,WHITE); pill(s,'DBMS · UNIT 3',1.05,1.84,1.35,LAV,PURPLE); text(s,'Question 1 of 3',7.2,1.87,1.25,.25,10,MUTED,align=PP_ALIGN.RIGHT)
text(s,'Which normal form removes partial dependency\non a composite candidate key?',1.05,2.45,6.95,.82,22,INK,True)
for i,opt in enumerate(['A. First Normal Form','B. Second Normal Form','C. Third Normal Form','D. Boyce–Codd Normal Form']):
    x=1.05+(i%2)*3.66; y=3.55+(i//2)*.76; rect(s,x,y,3.35,.58,MINT if i==1 else BG,line=GREEN if i==1 else RGBColor(225,229,237)); text(s,opt,x+.15,y,3.0,.58,11,GREEN if i==1 else INK,i==1)
rect(s,1.05,5.32,7.32,.7,LAV); text(s,'Why?  2NF removes partial dependencies of non-prime attributes.',1.25,5.43,6.9,.42,11,INK)
text(s,'SMART PRACTICE',9.45,1.65,2.3,.3,10,PURPLE,True)
items=[('Choose  ','DBMS, OS or Computer Networks',PURPLE),('Focus  ','Optional timed mode',ORANGE),('Learn  ','Answer explanations',GREEN),('Improve  ','Result history and score summaries',PURPLE)]
rich(s,items,9.45,2.18,3.0,3.65,14)
footer(s,6)

# 7 Responsible resources
s=base('Useful resources, handled responsibly','CONTENT PRINCIPLES')
rect(s,.7,1.55,5.6,4.8,NAVY); text(s,'“',1.05,1.82,.6,.6,38,RGBColor(156,143,245),True); text(s,'Important for preparation —\nnever guaranteed to appear.',1.1,2.55,4.55,1.1,27,WHITE,True); text(s,'Patterns help students prioritize.\nThey do not predict examinations.',1.1,4.1,4.4,.75,15,RGBColor(198,204,220)); pill(s,'HONEST GUIDANCE',1.1,5.38,1.55,RGBColor(49,59,103),RGBColor(215,211,255))
principles=[('✓','Original model papers','Created for practice.'),('✓','Authorized materials','Institution-provided or licensed.'),('✓','Visible licensing','Source and usage rights displayed.'),('✓','Protected uploads','Admin review before publishing.')]
for i,(ic,t,b) in enumerate(principles):
    y=1.65+i*1.18; icon_badge(s,ic,6.92,y,MINT,GREEN,.48); text(s,t,7.58,y-.02,4.3,.32,15,INK,True); text(s,b,7.58,y+.35,4.25,.3,12,MUTED)
footer(s,7)

# 8 Tech
s=base('A scalable, secure foundation','TECHNOLOGY')
tech=[('UI','Responsive HTML, CSS & JavaScript','Fast, dependency-light and accessible.',LAV,PURPLE),('DB','Supabase PostgreSQL','Normalized catalog and student data.',MINT,GREEN),('AU','Supabase Auth + RLS','Role-aware security at the database layer.',CORAL,ORANGE),('PWA','Service Worker','Installable app and offline resource access.',BLUE,RGBColor(49,112,187))]
for i,(ic,t,b,f,c) in enumerate(tech):
    x=.7+(i%2)*6.04; y=1.55+(i//2)*2.34; rect(s,x,y,5.64,1.95,WHITE); icon_badge(s,ic,x+.3,y+.36,f,c,.66); text(s,t,x+1.18,y+.29,3.95,.4,17,INK,True); text(s,b,x+1.18,y+.83,3.95,.58,13,MUTED)
footer(s,8)

# 9 Live product
s=base('Live, tested and ready to evolve','DEPLOYMENT')
links=[('VERCEL','Primary production app','backlog-buddy-tech-minds10.vercel.app',LAV,PURPLE),('RENDER','Secondary live deployment','backlog-buddy-kc5k.onrender.com',MINT,GREEN),('GITHUB','Source repository','github.com/sireeshareddykallam-beep/backlog-buddy',BLUE,RGBColor(49,112,187))]
for i,(lab,sub,url,f,c) in enumerate(links):
    y=1.55+i*1.45; rect(s,.7,y,11.92,1.12,WHITE); pill(s,lab,1.02,y+.36,1.15,f,c); text(s,sub,2.48,y+.18,3.3,.32,13,INK,True); text(s,url,2.48,y+.55,8.5,.3,12,PURPLE)
text(s,'Next opportunities',.7,6.08,2.2,.3,12,MUTED,True); text(s,'More universities  •  Authorized paper uploads  •  Richer analytics  •  Personalized recommendations',2.8,6.04,9.3,.38,13,INK)
footer(s,9)

# 10 close
s=base(dark=True)
text(s,'BACKLOG BUDDY',.75,.72,3.1,.35,12,RGBColor(184,176,255),True)
text(s,'A backlog is a chapter.\nNot the whole story.',.75,1.55,8.9,1.55,42,WHITE,True)
text(s,'Help students organize, practice and make their comeback.',.78,3.48,7.4,.45,20,RGBColor(198,204,220))
rect(s,.78,4.45,4.05,.68,PURPLE); text(s,'OPEN THE LIVE WEBSITE  →',.78,4.45,4.05,.68,12,WHITE,True,align=PP_ALIGN.CENTER)
text(s,'backlog-buddy-tech-minds10.vercel.app',.8,5.4,5.1,.32,13,RGBColor(184,176,255),True)
# decorative circles
for x,y,d,c in [(9.7,1.0,2.4,PURPLE),(10.75,2.65,1.35,RGBColor(55,69,120)),(8.92,3.82,2.05,RGBColor(75,62,163))]:
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=c;sh.line.color.rgb=c
text(s,'BB',9.66,1.0,2.4,2.4,34,WHITE,True,align=PP_ALIGN.CENTER)
footer(s,10,True)

prs.save(OUT)
print(OUT, OUT.stat().st_size)
